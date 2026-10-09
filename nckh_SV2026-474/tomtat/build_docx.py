import copy, re, sys
import docx
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TPL = "/home/nhanhuynh/Documents/pfdet_drone_follow/Mẫu tóm tắt đề tài.docx"
OUT = "/home/nhanhuynh/Documents/pfdet_drone_follow/nckh_SV2026-474/tomtat/tomtat.docx"
HINH = "/home/nhanhuynh/Documents/pfdet_drone_follow/nckh_SV2026-474/hinh/"

d = docx.Document(TPL)
body = d.element.body
els = list(body.iterchildren())

def set_text(p_el, text):
    """Giữ định dạng run đầu, thay chữ; hỗ trợ **đậm**."""
    runs = p_el.findall(qn('w:r'))
    proto = copy.deepcopy(runs[0]) if runs else None
    for r in p_el.findall(qn('w:r')) + p_el.findall(qn('w:hyperlink')):
        p_el.remove(r)
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if not part:
            continue
        bold = part.startswith('**')
        part = part.strip('*') if bold else part
        r = copy.deepcopy(proto)
        for t in r.findall(qn('w:t')) + r.findall(qn('w:drawing')) + r.findall(qn('w:br')):
            r.remove(t)
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = r.makeelement(qn('w:rPr'), {}); r.insert(0, rpr)
        for b in rpr.findall(qn('w:b')):
            rpr.remove(b)
        if bold:
            rpr.append(rpr.makeelement(qn('w:b'), {}))
        t = r.makeelement(qn('w:t'), {}); t.text = part
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t)
        p_el.append(r)

# ---- tiêu đề
set_text(els[0], "MÁY BAY KHÔNG NGƯỜI LÁI TỰ ĐỘNG THEO DÕI CON NGƯỜI BẰNG TRÍ TUỆ NHÂN TẠO")

# ---- bảng thông tin
tbl = d.tables[0]
vals = {
    0: "Huỳnh Lê Thành Nhân",
    1: "",
    2: "Trịnh Trần Quốc Bảo",
    3: "PGS.TS. Lê Mỹ Hà",
    4: "0937395656",
    5: "24119068@student.hcmute.edu.vn",
}
for i, row in enumerate(tbl.rows):
    if i in vals:
        c = row.cells[-1]
        p = c.paragraphs[0]._p
        set_text(p, vals[i])

# ảnh 3x4 của chủ nhiệm đề tài: thay ô placeholder ở cột đầu bằng ảnh thật
photo_cell = tbl.rows[0].cells[0]
for extra in photo_cell.paragraphs[1:]:
    extra._p.getparent().remove(extra._p)
pp = photo_cell.paragraphs[0]
for r in pp.runs:
    r._r.getparent().remove(r._r)
pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
pp.add_run().add_picture(HINH + "anh_3x4.jpg", width=Cm(3.0))

# ---- nguyên mẫu
P_HEAD, P_BODY, P_IMG, P_CAP, P_TCAP, T_TBL = (copy.deepcopy(els[k]) for k in (5, 6, 10, 11, 18, 19))
anchor = els[26]
for k in range(5, 26):
    body.remove(els[k])

def add(el):
    anchor.addprevious(el)

def head(t):
    e = copy.deepcopy(P_HEAD); set_text(e, "**"+t+"**"); add(e)

def para(t):
    e = copy.deepcopy(P_BODY); set_text(e, t); add(e)

def image(path, width_cm, cap):
    # chèn ảnh bằng python-docx vào đoạn tạm cuối tài liệu rồi chuyển run sang nguyên mẫu
    tmp = d.add_paragraph()
    run = tmp.add_run()
    run.add_picture(path, width=Cm(width_cm))
    e = copy.deepcopy(P_IMG)
    for r in e.findall(qn('w:r')):
        e.remove(r)
    e.append(run._r)
    tmp._p.getparent().remove(tmp._p)
    add(e)
    c = copy.deepcopy(P_CAP); set_text(c, cap); add(c)

def table(cap, rows):
    c = copy.deepcopy(P_TCAP); set_text(c, cap); add(c)
    t = copy.deepcopy(T_TBL)
    trs = t.findall(qn('w:tr'))
    proto_row = trs[1]
    for tr in trs:
        t.remove(tr)
    for ri, cells in enumerate(rows):
        tr = copy.deepcopy(trs[0] if ri == 0 else proto_row)
        tcs = tr.findall(qn('w:tc'))
        for ci, tc in enumerate(tcs):
            ps = tc.findall(qn('w:p'))
            for extra in ps[1:]:
                tc.remove(extra)
            txt = cells[ci] if ci < len(cells) else ""
            set_text(ps[0], txt)
        t.append(tr)
    add(t)

head("Mục tiêu đề tài")
para("Máy bay không người lái (UAV) được dùng ngày càng nhiều trong giám sát, cứu hộ và quay phim, nhưng người vận hành vừa phải lái vừa phải giữ đối tượng trong khung hình. Đề tài xây dựng một UAV tự phát hiện và bám theo người với ba mục tiêu: (i) mô hình trí tuệ nhân tạo nhận dạng người thời gian thực, độ chính xác từ 70% trở lên; (ii) thuật toán theo dõi giúp máy bay giữ khoảng cách và hướng bay phù hợp; (iii) tích hợp camera – xử lý ảnh – điều khiển bay thành hệ thống hoạt động ổn định ngoài thực tế. Toàn bộ xử lý chạy ngay trên máy bay, không cần kết nối Internet.")
head("Tính mới và tính sáng tạo")
para("Khung máy bay nhiều cánh quạt luôn rung khi bay, làm ảnh bị nhoè, nhiễu và dao động sáng. Bằng một giao thức đo thống nhất, cố định seed ngẫu nhiên và lặp ba lần, nhóm chỉ ra rằng mọi bộ phát hiện nhỏ gọn được thử (YOLOv8/v10/11/26 bản nano, RT-DETR-l) mất 74–83% độ chính xác khi rung. Đề tài đề xuất: (1) **mô hình SteadyDet-N** chỉ 1,17 triệu tham số, 4,6 GFLOPs, giữ nhánh độ phân giải cao cho người nhỏ nhưng làm mỏng để chạy trên thiết bị biên; (2) **huấn luyện nhận biết rung**: đưa mô hình vật lý của rung vào dữ liệu huấn luyện, không đổi nhãn, không tốn chi phí khi suy luận, nâng tỷ lệ giữ độ chính xác khi rung từ **16,8% lên 36,2%**; (3) chứng minh trên phần cứng thật rằng tốc độ trên NPU Hailo-8L phụ thuộc số context, SteadyDet-N nhanh hơn YOLO26-n **1,70 lần**; (4) lớp an toàn theo nguyên tắc người lái luôn giữ quyền chế độ bay, được khoá bằng kiểm thử tự động.")
head("Kết quả nghiên cứu")
para("Hệ thống gồm Raspberry Pi 5 gắn NPU Hailo-8L, camera USB và bộ điều khiển bay Pixhawk 6C chạy ArduPilot. Chuỗi xử lý: phát hiện người (SteadyDet-N INT8) → theo vết OC-SORT → khoá mục tiêu → ước lượng khoảng cách theo mô hình camera lỗ kim và bộ lọc Kalman → ba bộ PID → lớp an toàn → lệnh vận tốc MAVLink 30 Hz ở chế độ GUIDED.")
image(HINH + "vongbam_khung.png", 8.0, "Hình 1. Vòng bám người trên video bay thử ở độ cao khoảng 6 m")
para("Mô hình chạy trên máy bay đạt **AP@0,5 = 91,7%** trên dữ liệu bay thật, vượt chỉ tiêu 70% của thuyết minh. Khi rung, SteadyDet-N là mô hình nhỏ và rẻ nhất nhưng vượt cả bốn mô hình YOLO bản nano, và vượt RT-DETR-l (lớn gấp 28 lần) ở mức rung của bay tiến thông thường. Trên dữ liệu bay thật, xử lý rung làm độ khít hộp AP@0,75 tăng khoảng 80%.")
image(HINH + "duong_cong_rung.png", 8.0, "Hình 2. AP theo mức rung: sau huấn luyện nhận biết rung, SteadyDet-N vượt YOLO11-n và RT-DETR-l từ mức 2")
table("Bảng 1. Kết quả chính đo trên phần cứng thật", [
    ["Stt", "Chỉ số", "Kết quả"],
    ["1", "AP@0,5 trên dữ liệu bay thật (INT8)", "0,917"],
    ["2", "F1 / recall / precision", "0,880 / 0,837 / 0,928"],
    ["3", "Tốc độ đầu-cuối trên máy bay", "30,0 FPS"],
    ["4", "Độ trễ trung vị / p99", "32,04 / 36,15 ms"],
    ["5", "Suy luận NPU @320 / @640", "112,55 / 63,66 FPS"],
    ["6", "Tỷ lệ giữ độ chính xác khi rung", "16,8% → 36,2%"],
    ["7", "Khứ hồi lệnh MAVLink p50 / p99", "4,75 / 7,73 ms"],
    ["8", "Vòng bám trên video bay thử: giữ mục tiêu", "83,5% thời gian"],
    ["9", "Sai số hướng khi bám (trung vị)", "5,5°"],
])
para("Chương trình điều khiển có theo vết, khoá và tìm lại mục tiêu, thang an toàn khi mất mục tiêu và bảng giám sát web; hệ thống được kiểm thử qua bảy tầng từ kiểm thử tự động tới bay thử ngoài bãi. Chạy trên video bay thử, vòng bám khoá và giữ được người đi bộ trong 83,5% thời gian.")
head("Đóng góp về mặt kinh tế-xã hội, giáo dục và đào tạo, an ninh quốc phòng và khả năng áp dụng thực tế")
para("Hệ thống dùng linh kiện phổ thông, xử lý hoàn toàn trên máy bay nên hoạt động được ở vùng không có sóng di động, phù hợp tìm kiếm cứu nạn sau thiên tai, giám sát an ninh, tuần tra khu vực và quay phim tự động. Phương pháp huấn luyện nhận biết rung chỉ thay đổi dữ liệu huấn luyện, áp dụng được cho camera trên robot, xe tự hành mà không cần gimbal chống rung đắt tiền. Về đào tạo, đề tài cung cấp quy trình trọn vẹn, tái lập được từ dữ liệu, huấn luyện, lượng tử hoá, triển khai trên NPU tới điều khiển bay, dùng làm học liệu cho các học phần thị giác máy tính, hệ thống nhúng và điều khiển tự động.")
head("Công bố khoa học của SV từ kết quả nghiên cứu của đề tài")
para("Chưa có.")
# kết thúc tài liệu bằng chính section 2 cột: chuyển sectPr của đoạn ngắt vào cuối body
sp = anchor.find('.//'+qn('w:sectPr'))
fs = body.find(qn('w:sectPr'))
new = copy.deepcopy(sp)
body.replace(fs, new)
for el in list(body.iterchildren()):
    if el is new: continue
    if el.tag == qn('w:p') and el is not anchor and list(body.iterchildren()).index(el) > list(body.iterchildren()).index(anchor):
        body.remove(el)
body.remove(anchor)
d.save(OUT)
print("saved", OUT)
